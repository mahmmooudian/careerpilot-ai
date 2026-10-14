from io import BytesIO

import pymupdf
import pytest
from docx import Document

from app.services.resume_parser import ResumeParseError, parse_resume


def make_pdf(text: str) -> bytes:
    with pymupdf.open() as document:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
        return document.tobytes()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Data Scientist")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "2 years"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_resume_is_extracted():
    result = parse_resume("resume.PDF", make_pdf("Machine Learning Engineer with Python and SQL"))
    assert result.file_type == "pdf"
    assert result.page_count == 1
    assert "Machine Learning" in result.raw_text
    assert result.schema_version == "1.0"


def test_docx_table_content_is_extracted():
    result = parse_resume("resume.docx", make_docx())
    assert result.page_count is None
    assert "Python" in result.raw_text
    assert "Experience" in result.raw_text
    assert "2 years" in result.raw_text


def test_blank_pdf_warns_about_scans():
    result = parse_resume("blank.pdf", make_pdf(""))
    assert result.raw_text == ""
    assert any("OCR" in warning for warning in result.warnings)


def test_disguised_pdf_is_rejected():
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("fraud.pdf", b"not a PDF")
    assert raised.value.status_code == 422


def test_unsupported_extension_is_rejected():
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("resume.txt", b"text")
    assert raised.value.status_code == 415


def test_empty_upload_is_rejected():
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("resume.pdf", b"")
    assert raised.value.status_code == 400


def test_docx_header_is_extracted():
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Contact: hello@example.org"
    document.add_paragraph("Backend Engineer")
    buf = BytesIO()
    document.save(buf)
    result = parse_resume("resume.docx", buf.getvalue())
    assert "hello@example.org" in result.raw_text
    assert "Backend Engineer" in result.raw_text


def test_corrupt_docx_is_rejected():
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("resume.docx", b"this is not a ZIP file")
    assert raised.value.status_code == 422


def test_large_pdf_is_rejected():
    with pymupdf.open() as document:
        for _ in range(21):
            document.new_page()
        data = document.tobytes()
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("resume.pdf", data)
    assert raised.value.status_code == 413


def test_windows_filename_is_sanitized():
    result = parse_resume(r"C:\Users\someone\resume.docx", make_docx())
    assert result.file_name == "resume.docx"
